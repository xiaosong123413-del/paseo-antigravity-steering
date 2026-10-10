"""Small format-preservation fixtures, not user documents."""
from pathlib import Path
import sys


def create_fixtures(root):
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from pptx import Presentation
    from pptx.util import Inches
    import pymupdf
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph().add_run("ORIGINAL_1010").bold = True
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "KEEP_TABLE"
    doc.sections[0].header.paragraphs[0].text = "KEEP_HEADER"
    doc.save(root / "report.docx")
    book = Workbook()
    book.active["A1"] = "ORIGINAL_1010"
    book.active["A1"].font = Font(bold=True)
    book.active["B2"] = "=1+2"
    book.create_sheet("KeepSheet")["A1"] = "KEEP_SHEET"
    book.save(root / "report.xlsx")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "ORIGINAL_1010"
    run.font.bold = True
    slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1)).text = "KEEP_SHAPE"
    slides.save(root / "report.pptx")
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), "ORIGINAL_1010")
        page.insert_text((72, 144), "KEEP_PDF")
        pdf.save(root / "report.pdf")
    (root / "notes.md").write_bytes(b"# ORIGINAL_1010\r\n\r\nKEEP_MARKDOWN\r\n")


def edit_fixture(path, replacement="EDITED_1010"):
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation
    import pymupdf
    path = Path(path)
    if path.suffix == ".docx":
        doc = Document(path)
        doc.paragraphs[0].runs[0].text = replacement
        doc.save(path)
    elif path.suffix == ".xlsx":
        book = load_workbook(path)
        book.active["A1"] = replacement
        book.save(path)
        book.close()
    elif path.suffix == ".pptx":
        slides = Presentation(path)
        slides.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].text = replacement
        slides.save(path)
    elif path.suffix == ".pdf":
        with pymupdf.open(path) as pdf:
            page = pdf[0]
            for rect in page.search_for("ORIGINAL_1010"):
                page.add_redact_annot(rect)
            page.apply_redactions()
            page.insert_text((72, 72), replacement)
            pdf.saveIncr()
    else:
        path.write_bytes(path.read_bytes().replace(b"ORIGINAL_1010", replacement.encode()))


def assert_fixture(path, expected="EDITED_1010"):
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation
    import pymupdf
    path = Path(path)
    if path.suffix == ".docx":
        doc = Document(path)
        assert doc.paragraphs[0].text == expected
        assert doc.paragraphs[0].runs[0].bold
        assert doc.tables[0].cell(0, 0).text == "KEEP_TABLE"
        assert doc.sections[0].header.paragraphs[0].text == "KEEP_HEADER"
    elif path.suffix == ".xlsx":
        book = load_workbook(path, data_only=False)
        assert book.active["A1"].value == expected
        assert book.active["A1"].font.bold
        assert book.active["B2"].value == "=1+2"
        assert book["KeepSheet"]["A1"].value == "KEEP_SHEET"
        book.close()
    elif path.suffix == ".pptx":
        slides = Presentation(path)
        assert len(slides.slides) == 1
        assert slides.slides[0].shapes[0].text == expected
        assert slides.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.bold
        assert slides.slides[0].shapes[1].text == "KEEP_SHAPE"
    elif path.suffix == ".pdf":
        with pymupdf.open(path) as pdf:
            assert pdf.page_count == 1
            text = pdf[0].get_text()
            assert expected in text and "KEEP_PDF" in text
            assert "ORIGINAL_1010" not in text or expected == "ORIGINAL_1010"
    else:
        data = path.read_bytes()
        assert expected.encode() in data and b"KEEP_MARKDOWN" in data
        assert b"\r\n" in data


if __name__ == "__main__":
    create_fixtures(sys.argv[1])
    print("Created five document fixtures")
