"""Stage, validate, and publish document edits without flattening the source format."""

import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import zipfile

SUPPORTED = {".docx", ".xlsx", ".pptx", ".pdf", ".md", ".markdown", ".txt", ".csv"}


def fingerprint(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def validate(path):
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported format {suffix}. Use DOCX/XLSX/PPTX/PDF/Markdown/TXT/CSV; legacy Office files require conversion first.")
    if suffix in {".docx", ".xlsx", ".pptx"}:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            if broken:
                raise ValueError(f"Corrupt archive member: {broken}")
    if suffix == ".docx":
        from docx import Document
        doc = Document(path)
        return {"format": suffix, "paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            return {"format": suffix, "sheets": book.sheetnames}
        finally:
            book.close()
    if suffix == ".pptx":
        from pptx import Presentation
        return {"format": suffix, "slides": len(Presentation(path).slides)}
    if suffix == ".pdf":
        import pymupdf
        with pymupdf.open(path) as doc:
            if doc.needs_pass:
                raise ValueError("Encrypted PDF needs a password before editing")
            return {"format": suffix, "pages": doc.page_count}
    data = path.read_bytes()
    try:
        data.decode("utf-8-sig")
        encoding = "utf-8"
    except UnicodeDecodeError:
        encoding = "other; detect encoding before editing"
    return {"format": suffix, "bytes": len(data), "encoding": encoding}


def prepare(source, output_root, name=None, uploaded=False):
    source = Path(source).resolve(strict=True)
    if not source.is_file():
        raise ValueError("Source must be a regular file")
    file_name = name or source.name
    if Path(file_name).name != file_name or any(c in file_name for c in '\\/:<>"|?*'):
        raise ValueError("Name must be a file name, not a path")
    if Path(file_name).suffix.lower() not in SUPPORTED:
        raise ValueError("Unsupported document extension; legacy .doc/.xls/.ppt require conversion first")
    root = Path(output_root).resolve()
    root.mkdir(parents=True, exist_ok=True)
    job = Path(tempfile.mkdtemp(prefix="edit-", dir=root))
    working = job / file_name
    original_hash = fingerprint(source)
    shutil.copy2(source, working)
    if fingerprint(working) != original_hash or fingerprint(source) != original_hash:
        raise ValueError("Source changed during staging; prepare again")
    details = validate(working)
    manifest = job / "edit.json"
    manifest.write_text(json.dumps({"version": 1, "source": str(source),
        "sha256": original_hash, "working": str(working), "uploaded": uploaded}, ensure_ascii=False), encoding="utf-8")
    return {"manifest": str(manifest), "working": str(working), "source": str(source),
            "default": "new file", "validation": details}


def publish(manifest_path, overwrite_original=False):
    manifest_path = Path(manifest_path).resolve(strict=True)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("version") != 1:
        raise ValueError("Unknown edit manifest version")
    source = Path(manifest["source"]).resolve(strict=True)
    working = Path(manifest["working"]).resolve(strict=True)
    if working.parent != manifest_path.parent or source == working:
        raise ValueError("Working file must be inside its edit directory and separate from the source")
    details = validate(working)
    if fingerprint(source) != manifest["sha256"]:
        raise ValueError("Original changed since preparation; refusing to overwrite or claim it is unchanged")
    result = working
    if overwrite_original:
        if manifest["uploaded"]:
            raise ValueError("An upload is a copy, not the original on the user's computer. Prepare the explicit local source path instead.")
        if source.suffix.lower() != working.suffix.lower():
            raise ValueError("Cannot overwrite the source with a different format")
        fd, temporary = tempfile.mkstemp(prefix=".paseo-edit-", suffix=source.suffix, dir=source.parent)
        os.close(fd)
        temporary = Path(temporary)
        try:
            shutil.copy2(working, temporary)
            validate(temporary)
            if fingerprint(source) != manifest["sha256"]:
                raise ValueError("Original changed before commit; refusing overwrite")
            os.replace(temporary, source)
        finally:
            temporary.unlink(missing_ok=True)
        result = source
    return {"path": str(result), "mode": "original" if overwrite_original else "new_file",
            "sha256": fingerprint(result), "validation": details,
            "link": f"[{result.name}](<{result.as_posix()}>)"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    stage = commands.add_parser("prepare")
    stage.add_argument("source")
    stage.add_argument("--output-root", required=True)
    stage.add_argument("--name")
    stage.add_argument("--uploaded", action="store_true")
    commit = commands.add_parser("publish")
    commit.add_argument("manifest")
    commit.add_argument("--overwrite-original", action="store_true")
    check = commands.add_parser("validate")
    check.add_argument("path")
    args = parser.parse_args()
    try:
        if args.command == "prepare":
            result = prepare(args.source, args.output_root, args.name, args.uploaded)
        elif args.command == "publish":
            result = publish(args.manifest, args.overwrite_original)
        else:
            result = validate(Path(args.path))
        print(json.dumps(result, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False), file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
